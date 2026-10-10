import os
import sys
import gc
import json
import re
import hashlib
import time
import urllib.request
import urllib.error
from difflib import SequenceMatcher
from datetime import datetime, timedelta
from pathlib import Path

AI_ONLY = "--ai-only" in sys.argv

# =========================================================
# Windows + CUDA DLL 경로 설정
# - pip로 설치한 nvidia-cublas-cu12 / nvidia-cudnn-cu12 사용
# - os.add_dll_directory() 반환 핸들은 반드시 유지해야 함
# =========================================================

dll_handles = []

if os.name == "nt":
    site_packages = Path(sys.prefix) / "Lib" / "site-packages"
    cublas_bin = site_packages / "nvidia" / "cublas" / "bin"
    cudnn_bin = site_packages / "nvidia" / "cudnn" / "bin"

    for dll_dir in (cublas_bin, cudnn_bin):
        if dll_dir.exists():
            try:
                dll_handles.append(os.add_dll_directory(str(dll_dir)))
            except (AttributeError, OSError):
                pass
            os.environ["PATH"] = str(dll_dir) + os.pathsep + os.environ.get("PATH", "")

from faster_whisper import WhisperModel
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


# =========================================================
# 설정
# =========================================================

CALL_DIR = Path(r"E:\통화녹음")

# RTX 4050 6GB 기준: 속도/정확도 균형
WHISPER_MODEL = "medium"
WHISPER_DEVICE = "cuda"
WHISPER_COMPUTE_TYPE = "float16"

# 1.5B보다 업무 문맥/요약 품질 개선 목적
OLLAMA_MODEL = "qwen2.5:3b"
OLLAMA_URL = "http://localhost:11434/api/chat"
OLLAMA_TIMEOUT = 180       # 1회 요청 최대 대기(초)
OLLAMA_RETRIES = 2         # 실패 시 추가 재시도 횟수
OLLAMA_NUM_CTX = 8192
# OLLAMA_NUM_PREDICT = 1200
OLLAMA_NUM_PREDICT = 2400

# 10분 내외 통화까지 안전하게 처리
# MAX_CHARS_PER_CHUNK = 7000
MAX_CHARS_PER_CHUNK = 4500
CHUNK_OVERLAP_LINES = 3

# 낮은 신뢰도는 삭제하지 않고 Excel에서 검토 표시
REVIEW_CONFIDENCE = 0.70

OUTPUT_FILE = CALL_DIR / "통화문의내역_개선.xlsx"
CACHE_DIR = CALL_DIR / "_분석캐시"
# TRANSCRIPT_DIR = CACHE_DIR / "녹취록"
TRANSCRIPT_DIR = CACHE_DIR / f"녹취록_{WHISPER_MODEL}"

# 프롬프트/모델을 바꿨을 때 과거 저품질 분석 캐시를 섞지 않기 위한 버전
# ANALYSIS_VERSION = "improved_v2_qwen25_3b"
ANALYSIS_VERSION = f"improved_v2_qwen25_3b_{WHISPER_MODEL}"
ANALYSIS_DIR = CACHE_DIR / ANALYSIS_VERSION

# 중간 결과 저장 주기
CHECKPOINT_EVERY = 10

AUDIO_EXTENSIONS = {
    ".m4a", ".mp3", ".wav", ".aac",
    ".flac", ".ogg", ".mp4", ".3gp"
}

VALID_CATEGORIES = {
    "장애/오류",
    "성능",
    "기능문의",
    "기능수정",
    "출력/문서",
    "데이터",
    "계정/권한",
    "파일/다운로드",
    "배포",
    "사용방법",
    "기타",
}

TRANSCRIPT_DIR.mkdir(parents=True, exist_ok=True)
ANALYSIS_DIR.mkdir(parents=True, exist_ok=True)


# =========================================================
# 공통 유틸
# =========================================================

def format_duration(seconds):
    if seconds < 60:
        return f"{seconds:.2f}초"

    total_seconds = int(seconds)
    hours = total_seconds // 3600
    minutes = (total_seconds % 3600) // 60
    secs = total_seconds % 60

    if hours > 0:
        return f"{hours}시간 {minutes}분 {secs}초"
    return f"{minutes}분 {secs}초"


def clean_text(value):
    if value is None:
        return ""
    value = str(value).strip()
    value = re.sub(r"\s+", " ", value)
    return value


def normalize_for_compare(value):
    value = clean_text(value).lower()
    value = re.sub(r"[^0-9a-z가-힣]+", "", value)
    return value


def text_similarity(a, b):
    a = normalize_for_compare(a)
    b = normalize_for_compare(b)
    if not a or not b:
        return 0.0
    return SequenceMatcher(None, a, b).ratio()


def parse_confidence(value):
    try:
        result = float(value)
        return max(0.0, min(1.0, result))
    except (TypeError, ValueError):
        return 0.5


def get_cache_key(path: Path):
    # 경로 + 파일 크기 + 수정시간을 포함해서 파일 교체 시 캐시가 자동 무효화되도록 함
    stat = path.stat()
    value = f"{path.resolve()}|{stat.st_size}|{stat.st_mtime_ns}".encode("utf-8")
    return hashlib.sha1(value).hexdigest()


# =========================================================
# 파일명 파싱
# =========================================================

def parse_filename(path: Path):
    """
    예:
    통화 녹음 디앤디기술 백현욱차장님_260417_141615.m4a
    통화 디앤디기술 백현욱차장님_261007_091215.m4a
    """
    stem = path.stem
    pattern = r"^(?:통화(?:\s*녹음)?\s*)?(.*?)_(\d{6}|\d{8})_(\d{6})$"
    match = re.match(pattern, stem)

    if match:
        person = match.group(1).strip()
        date_string = match.group(2)
        time_string = match.group(3)

        if len(date_string) == 6:
            date_obj = datetime.strptime(date_string, "%y%m%d")
        else:
            date_obj = datetime.strptime(date_string, "%Y%m%d")

        time_obj = datetime.strptime(time_string, "%H%M%S")

        return {
            "person": person,
            "date": date_obj.strftime("%Y-%m-%d"),
            "time": time_obj.strftime("%H:%M:%S"),
        }

    modified = datetime.fromtimestamp(path.stat().st_mtime)
    return {
        "person": stem,
        "date": modified.strftime("%Y-%m-%d"),
        "time": modified.strftime("%H:%M:%S"),
    }


# =========================================================
# STT
# =========================================================

def transcribe_audio(path: Path, model: WhisperModel):
    cache_key = get_cache_key(path)
    transcript_file = TRANSCRIPT_DIR / f"{cache_key}.txt"

    if transcript_file.exists():
        return transcript_file.read_text(encoding="utf-8"), True

    segments, _ = model.transcribe(
        str(path),
        language="ko",
        beam_size=3,
        vad_filter=True,
        condition_on_previous_text=False,
        vad_parameters={
            "min_silence_duration_ms": 500,
        },
    )

    lines = []
    for segment in segments:
        text = segment.text.strip()
        if text:
            lines.append(text)

    transcript = "\n".join(lines)
    transcript_file.write_text(transcript, encoding="utf-8")
    return transcript, False


# =========================================================
# Ollama
# =========================================================

SYSTEM_PROMPT = r"""
당신은 IT 시스템 유지보수 업무의 실제 통화 기록을 정리하는 담당자입니다.
목표는 '업무 문의/요청/장애/수정사항'을 엑셀 한 행 단위로 정확하고 간결하게 정리하는 것입니다.

반드시 아래 원칙을 지키세요.

[추출 원칙]
1. 녹취에 실제로 존재하는 내용만 사용합니다. 없는 원인, 해결책, 기능, 담당자 의견을 절대 만들어내지 않습니다.
2. 인사, 안부, 식사, 위치 확인, 단순 통화 연결, 잡담 등 업무 이슈가 아닌 내용은 제외합니다.
3. 하나의 통화에서 같은 주제에 대해 여러 번 대화했다면 하나의 문의로 합칩니다.
4. 서로 독립적인 업무 이슈가 실제로 여러 개라면 여러 문의로 분리합니다.
5. '문의내용', '답변/처리내용', '후속조치'를 서로 구분합니다.
6. 상대방이 문의한 문장을 답변내용에 그대로 복사하지 않습니다.
7. 답변이 녹취에 명확히 없으면 response는 빈 문자열로 둡니다.
8. 후속조치가 명확히 없으면 follow_up은 빈 문자열로 둡니다.
9. 녹취가 불명확해도 추측해서 채우지 말고 confidence를 낮게 줍니다.
10. 일반적인 기술 상식으로 해결책을 추가하지 않습니다. 예: '캐시 삭제', '자바스크립트 최적화' 등이 실제 통화에서 언급되지 않았다면 절대 쓰지 않습니다.
11. title은 엑셀에서 한눈에 볼 수 있도록 핵심만 15~35자 정도로 씁니다.
12. inquiry는 무엇이 안 되는지/무엇을 요청했는지가 드러나게 씁니다.
13. summary는 '통화 전체 요약' 같은 문구가 아니라 실제 통화 내용을 1~2문장으로 요약합니다.
14. 업무 이슈가 전혀 없으면 inquiries를 빈 배열로 반환합니다.

[상태 판단]
- 완료: 통화 중 조치가 끝났거나 해결됐다고 명확히 확인됨
- 처리예정: 수정, 배포, 확인, 회신 등 이후 작업이 명확히 남아 있음
- 확인필요: 결론이 불명확하거나 추가 확인이 필요함

[분류]
아래 중 하나만 사용합니다.
- 장애/오류
- 성능
- 기능문의
- 기능수정
- 출력/문서
- 데이터
- 계정/권한
- 파일/다운로드
- 배포
- 사용방법
- 기타

[evidence]
- 해당 문의를 판단하는 데 직접 근거가 된 녹취 표현을 짧게 적습니다.
- 녹취에 없는 문장을 새로 만들지 않습니다.

[confidence]
- 0.0~1.0 숫자
- 0.9 이상: 매우 명확
- 0.7~0.89: 대체로 명확
- 0.5~0.69: 일부 불명확
- 0.5 미만: 검토 필요

반드시 JSON만 반환하세요.

형식:
{
  "summary": "실제 통화 요약",
  "inquiries": [
    {
      "is_valid": true,
      "category": "분류",
      "title": "문의 제목",
      "inquiry": "문의 내용",
      "response": "답변 또는 처리 내용",
      "follow_up": "후속 조치",
      "status": "완료 또는 처리예정 또는 확인필요",
      "evidence": "녹취 근거 표현",
      "confidence": 0.0
    }
  ]
}
"""


def check_ollama():
    try:
        with urllib.request.urlopen("http://localhost:11434/api/tags", timeout=5) as response:
            return response.status == 200
    except Exception:
        return False


def call_ollama(prompt, timeout=OLLAMA_TIMEOUT, retries=OLLAMA_RETRIES):

    last_error = None

    for attempt in range(retries + 1):

        if attempt == 0:
            current_prompt = prompt
        else:
            current_prompt = (
                    "이전 AI 응답은 문장을 반복하다 실패했습니다.\n"
                    "이번에는 다음 규칙을 반드시 지키세요.\n"
                    "1. 같은 단어나 문장을 반복하지 마세요.\n"
                    "2. 인사말과 의미 없는 반복 대화는 제외하세요.\n"
                    "3. 문의내용은 항목당 최대 100자로 요약하세요.\n"
                    "4. 답변과 후속조치도 각각 최대 100자로 작성하세요.\n"
                    "5. 통화 전체 내용을 복사하지 마세요.\n"
                    "6. 업무 문의가 없다면 inquiries는 []로 반환하세요.\n"
                    "7. 반드시 완결된 JSON을 반환하세요.\n\n"
                    + prompt
            )

        body = {
            "model": OLLAMA_MODEL,
            "stream": False,
            "format": "json",
            "keep_alive": "10m",
            "messages": [
                {
                    "role": "system",
                    "content": SYSTEM_PROMPT
                },
                {
                    "role": "user",
                    "content": current_prompt
                }
            ],
            "options": {
                "temperature": 0.0,
                "num_ctx": OLLAMA_NUM_CTX,
                "num_predict": (
                    OLLAMA_NUM_PREDICT
                    if attempt == 0
                    else 1200
                ),
                "repeat_penalty": 1.15
            }
        }

        data = json.dumps(
            body,
            ensure_ascii=False
        ).encode("utf-8")

        request = urllib.request.Request(
            OLLAMA_URL,
            data=data,
            headers={
                "Content-Type": "application/json"
            }
        )

        try:
            with urllib.request.urlopen(
                    request,
                    timeout=timeout
            ) as response:
                result = json.loads(
                    response.read().decode("utf-8")
                )

            content = result["message"]["content"]

            try:
                return json.loads(content)

            except json.JSONDecodeError as error:

                print(f"      JSON 오류: {error}")
                print(
                    f"      종료 사유: "
                    f"{result.get('done_reason')}"
                )
                print(
                    f"      출력 토큰: "
                    f"{result.get('eval_count')}"
                )
                print(
                    f"      마지막 내용: "
                    f"{repr(content[-200:])}"
                )

                raise

        except (
                TimeoutError,
                urllib.error.URLError,
                json.JSONDecodeError,
                KeyError
        ) as error:

            last_error = error

            if attempt >= retries:
                break

            wait = 3 * (attempt + 1)

            print(
                f"      AI 호출 실패({error})"
                f" - {wait}초 후 재시도"
            )

            time.sleep(wait)

    raise RuntimeError(
        f"Ollama 분석 실패: {last_error}"
    )

# =========================================================
# 긴 통화 분할 / AI 분석
# =========================================================

def split_transcript(text, max_chars=MAX_CHARS_PER_CHUNK):
    if len(text) <= max_chars:
        return [text]

    lines = text.splitlines()
    chunks = []
    current = []
    current_length = 0

    for line in lines:
        line_len = len(line) + 1

        if current and current_length + line_len > max_chars:
            chunks.append("\n".join(current))
            current = current[-CHUNK_OVERLAP_LINES:]
            current_length = sum(len(x) + 1 for x in current)

        current.append(line)
        current_length += line_len

    if current:
        chunks.append("\n".join(current))

    return chunks


def clean_ai_item(item):
    if not isinstance(item, dict):
        return None

    is_valid = item.get("is_valid", True)
    if isinstance(is_valid, str):
        is_valid = is_valid.strip().lower() not in {"false", "0", "no", "n"}
    if not is_valid:
        return None

    category = clean_text(item.get("category"))
    if category not in VALID_CATEGORIES:
        category = "기타"

    title = clean_text(item.get("title"))
    inquiry = clean_text(item.get("inquiry"))
    response = clean_text(item.get("response"))
    follow_up = clean_text(item.get("follow_up"))
    status = clean_text(item.get("status"))
    evidence = clean_text(item.get("evidence"))
    confidence = parse_confidence(item.get("confidence"))

    if status not in {"완료", "처리예정", "확인필요"}:
        status = "확인필요"

    # 실질적인 문의 내용이 없으면 제거
    if not inquiry or len(normalize_for_compare(inquiry)) < 4:
        return None

    # 문의를 답변/후속조치에 그대로 복제한 경우 제거
    if text_similarity(inquiry, response) >= 0.93:
        response = ""
    if text_similarity(inquiry, follow_up) >= 0.93 or text_similarity(response, follow_up) >= 0.95:
        follow_up = ""

    # 제목이 비어 있으면 문의 앞부분 사용
    if not title:
        title = inquiry[:35]

    return {
        "category": category,
        "title": title,
        "inquiry": inquiry,
        "response": response,
        "follow_up": follow_up,
        "status": status,
        "evidence": evidence,
        "confidence": confidence,
        "review": "Y" if confidence < REVIEW_CONFIDENCE else "",
    }


def merge_duplicate_inquiries(items):
    merged = []

    for item in items:
        duplicate_index = None

        for idx, existing in enumerate(merged):
            title_sim = text_similarity(item["title"], existing["title"])
            inquiry_sim = text_similarity(item["inquiry"], existing["inquiry"])

            if title_sim >= 0.86 or inquiry_sim >= 0.82:
                duplicate_index = idx
                break

        if duplicate_index is None:
            merged.append(item)
            continue

        existing = merged[duplicate_index]

        # 더 구체적인/긴 설명을 유지
        if len(item["inquiry"]) > len(existing["inquiry"]):
            existing["inquiry"] = item["inquiry"]
        if len(item["response"]) > len(existing["response"]):
            existing["response"] = item["response"]
        if len(item["follow_up"]) > len(existing["follow_up"]):
            existing["follow_up"] = item["follow_up"]
        if len(item["evidence"]) > len(existing["evidence"]):
            existing["evidence"] = item["evidence"]

        existing["confidence"] = max(existing["confidence"], item["confidence"])
        existing["review"] = "Y" if existing["confidence"] < REVIEW_CONFIDENCE else ""

        # 처리예정/확인필요 정보가 있으면 보수적으로 유지
        status_priority = {"완료": 1, "처리예정": 2, "확인필요": 3}
        if status_priority.get(item["status"], 3) > status_priority.get(existing["status"], 3):
            existing["status"] = item["status"]

    return merged


def consolidate_chunks(path, summaries, inquiries):
    """긴 통화가 여러 chunk로 나뉜 경우 마지막으로 중복 이슈를 정리한다."""
    if len(summaries) <= 1:
        return None

    payload = {
        "chunk_summaries": summaries,
        "inquiries": inquiries,
    }

    prompt = f"""
아래는 같은 통화를 여러 구간으로 나누어 분석한 임시 결과입니다.
같은 업무 이슈가 여러 번 나온 것은 하나로 합치고, 서로 다른 실제 업무 이슈만 남기세요.
녹취에 없는 내용을 추가하지 마세요.

임시 결과:
{json.dumps(payload, ensure_ascii=False, indent=2)}
"""

    try:
        return call_ollama(prompt)
    except Exception as error:
        print(f"      긴 통화 통합 분석 실패 - 기존 chunk 결과 사용: {error}")
        return None


def analyze_transcript(path, transcript):
    cache_key = get_cache_key(path)
    analysis_file = ANALYSIS_DIR / f"{cache_key}.json"

    if analysis_file.exists():
        return json.loads(analysis_file.read_text(encoding="utf-8")), True

    chunks = split_transcript(transcript)
    all_raw_inquiries = []
    summaries = []

    for index, chunk in enumerate(chunks, start=1):
        print(f"    AI 분석 중... ({index}/{len(chunks)})")

        prompt = f"""
아래는 실제 업무 통화의 STT 녹취입니다.
STT 특성상 일부 단어가 잘못 인식될 수 있으므로 전체 문맥을 기준으로 판단하되,
확실하지 않은 고유명사나 기술용어를 임의로 고쳐 쓰지 마세요.

[녹취 시작]
{chunk}
[녹취 끝]
"""

        result = call_ollama(prompt)

        summary = clean_text(result.get("summary"))
        if summary and summary not in {"통화 전체 요약", "전체 요약", "요약"}:
            summaries.append(summary)

        raw_items = result.get("inquiries", [])
        if isinstance(raw_items, list):
            all_raw_inquiries.extend(raw_items)

    # 긴 통화는 chunk 결과를 한 번 더 통합
    consolidated = consolidate_chunks(path, summaries, all_raw_inquiries)
    if consolidated:
        summary = clean_text(consolidated.get("summary"))
        if summary and summary not in {"통화 전체 요약", "전체 요약", "요약"}:
            summaries = [summary]
        raw_items = consolidated.get("inquiries", [])
        if isinstance(raw_items, list):
            all_raw_inquiries = raw_items

    cleaned = []
    for item in all_raw_inquiries:
        clean_item = clean_ai_item(item)
        if clean_item:
            cleaned.append(clean_item)

    cleaned = merge_duplicate_inquiries(cleaned)

    final_result = {
        "summary": " / ".join(dict.fromkeys(summaries)),
        "inquiries": cleaned,
    }

    analysis_file.write_text(
        json.dumps(final_result, ensure_ascii=False, indent=2),
        encoding="utf-8",
    )

    return final_result, False


# =========================================================
# Excel
# =========================================================

def style_sheet(sheet):
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions

    header_fill = PatternFill("solid", fgColor="D9EAF7")
    review_fill = PatternFill("solid", fgColor="FFF2CC")

    for cell in sheet[1]:
        cell.font = Font(bold=True)
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center")

    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)

        # '검토필요' 컬럼이 있으면 Y인 행 강조
        headers = {c.value: c.column for c in sheet[1]}
        review_col = headers.get("검토필요")
        if review_col and sheet.cell(row=row[0].row, column=review_col).value == "Y":
            for cell in row:
                cell.fill = review_fill


def set_widths(sheet, widths):
    for col, width in widths.items():
        sheet.column_dimensions[get_column_letter(col)].width = width


def create_excel(call_results, errors, output_path=OUTPUT_FILE):
    wb = Workbook()

    # 문의내역
    ws = wb.active
    ws.title = "문의내역"

    headers = [
        "번호",
        "통화일자",
        "통화시간",
        "상대방",
        "분류",
        "문의제목",
        "문의내용",
        "답변/처리내용",
        "후속조치",
        "상태",
        "신뢰도",
        "검토필요",
        "근거문구",
        "통화요약",
        "원본파일",
    ]
    ws.append(headers)

    row_number = 1
    for call in call_results:
        metadata = call["metadata"]
        analysis = call["analysis"]

        for inquiry in analysis.get("inquiries", []):
            ws.append([
                row_number,
                metadata["date"],
                metadata["time"],
                metadata["person"],
                inquiry.get("category", ""),
                inquiry.get("title", ""),
                inquiry.get("inquiry", ""),
                inquiry.get("response", ""),
                inquiry.get("follow_up", ""),
                inquiry.get("status", ""),
                round(float(inquiry.get("confidence", 0.0)), 2),
                inquiry.get("review", ""),
                inquiry.get("evidence", ""),
                analysis.get("summary", ""),
                call["filename"],
            ])
            row_number += 1

    # 통화목록
    ws_calls = wb.create_sheet("통화목록")
    ws_calls.append([
        "번호",
        "통화일자",
        "통화시간",
        "상대방",
        "문의건수",
        "검토필요건수",
        "통화요약",
        "원본파일",
    ])

    for index, call in enumerate(call_results, start=1):
        analysis = call["analysis"]
        inquiries = analysis.get("inquiries", [])
        review_count = sum(1 for item in inquiries if item.get("review") == "Y")
        metadata = call["metadata"]

        ws_calls.append([
            index,
            metadata["date"],
            metadata["time"],
            metadata["person"],
            len(inquiries),
            review_count,
            analysis.get("summary", ""),
            call["filename"],
        ])

    # 검토필요
    ws_review = wb.create_sheet("검토필요")
    ws_review.append(headers)
    review_no = 1

    for call in call_results:
        metadata = call["metadata"]
        analysis = call["analysis"]
        for inquiry in analysis.get("inquiries", []):
            if inquiry.get("review") != "Y":
                continue

            ws_review.append([
                review_no,
                metadata["date"],
                metadata["time"],
                metadata["person"],
                inquiry.get("category", ""),
                inquiry.get("title", ""),
                inquiry.get("inquiry", ""),
                inquiry.get("response", ""),
                inquiry.get("follow_up", ""),
                inquiry.get("status", ""),
                round(float(inquiry.get("confidence", 0.0)), 2),
                inquiry.get("review", ""),
                inquiry.get("evidence", ""),
                analysis.get("summary", ""),
                call["filename"],
            ])
            review_no += 1

    # 오류목록
    ws_errors = wb.create_sheet("오류목록")
    ws_errors.append(["번호", "원본파일", "단계", "오류"])
    for index, item in enumerate(errors, start=1):
        ws_errors.append([
            index,
            item.get("filename", ""),
            item.get("stage", ""),
            item.get("error", ""),
        ])

    for sheet in [ws, ws_calls, ws_review, ws_errors]:
        style_sheet(sheet)

    set_widths(ws, {
        1: 7, 2: 13, 3: 11, 4: 28, 5: 14,
        6: 32, 7: 58, 8: 58, 9: 48, 10: 13,
        11: 10, 12: 10, 13: 55, 14: 65, 15: 48,
    })
    set_widths(ws_calls, {
        1: 7, 2: 13, 3: 11, 4: 28, 5: 10,
        6: 12, 7: 75, 8: 48,
    })
    set_widths(ws_review, {
        1: 7, 2: 13, 3: 11, 4: 28, 5: 14,
        6: 32, 7: 58, 8: 58, 9: 48, 10: 13,
        11: 10, 12: 10, 13: 55, 14: 65, 15: 48,
    })
    set_widths(ws_errors, {1: 7, 2: 50, 3: 15, 4: 90})

    wb.save(output_path)


# =========================================================
# 진행률
# =========================================================

def print_eta(start_time, completed, total, label="전체"):
    if completed <= 0:
        return

    elapsed = time.perf_counter() - start_time
    avg = elapsed / completed
    remaining_count = max(0, total - completed)
    remaining = avg * remaining_count
    end_time = datetime.now() + timedelta(seconds=remaining)

    print(
        f"    {label} 평균 {format_duration(avg)}/건 | "
        f"남은 {remaining_count}건 | "
        f"예상 {format_duration(remaining)} | "
        f"종료 {end_time.strftime('%H:%M:%S')}"
    )


# =========================================================
# MAIN
# =========================================================

def main():
    print("=" * 72)
    print("갤럭시 통화녹음 업무문의 자동 분석 - 개선 버전")
    print("=" * 72)
    print(f"시작시간: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}")
    print(f"Whisper: {WHISPER_MODEL} / {WHISPER_DEVICE} / {WHISPER_COMPUTE_TYPE}")
    print(f"Ollama : {OLLAMA_MODEL} / ctx={OLLAMA_NUM_CTX}")
    print()

    if not CALL_DIR.exists():
        print(f"폴더가 없습니다: {CALL_DIR}")
        return

    if not check_ollama():
        print("Ollama에 연결할 수 없습니다.")
        print(f"먼저 실행/설치 확인: ollama pull {OLLAMA_MODEL}")
        return

    audio_files = sorted(
        file for file in CALL_DIR.iterdir()
        if file.is_file() and file.suffix.lower() in AUDIO_EXTENSIONS
    )

    total_files = len(audio_files)
    print(f"발견된 통화녹음: {total_files}개")

    if not audio_files:
        return

    errors = []
    transcripts = []

    # -----------------------------------------------------
    # 1단계: STT
    # - 이미 캐시가 있으면 Whisper를 아예 로딩하지 않음
    # - 새 STT가 발생한 경우 Windows/CUDA 해제 충돌을 피하기 위해
    #   현재 프로세스를 AI 전용 모드로 재시작함
    # -----------------------------------------------------

    missing_stt = []
    for path in audio_files:
        transcript_file = TRANSCRIPT_DIR / f"{get_cache_key(path)}.txt"
        if not transcript_file.exists():
            missing_stt.append(path)

    if AI_ONLY:
        print("\n[1단계] STT 건너뜀 (--ai-only)")

    elif not missing_stt:
        print("\n[1단계] STT 캐시 100% 존재 - Whisper 로딩 생략")

    else:
        print(f"\n[1단계] STT 처리 - 미처리 {len(missing_stt)}건")
        print("Whisper 모델 로딩 중...")
        model_start = time.perf_counter()

        whisper = WhisperModel(
            WHISPER_MODEL,
            device=WHISPER_DEVICE,
            compute_type=WHISPER_COMPUTE_TYPE,
        )

        print(f"Whisper 로딩 완료: {format_duration(time.perf_counter() - model_start)}")
        stt_phase_start = time.perf_counter()

        for index, path in enumerate(missing_stt, start=1):
            file_start = time.perf_counter()
            print(f"[{index}/{len(missing_stt)}] {path.name}")

            try:
                transcript, cached = transcribe_audio(path, whisper)
                elapsed = time.perf_counter() - file_start
                print(f"    STT 완료: {format_duration(elapsed)}")

                if not transcript.strip():
                    errors.append({
                        "filename": path.name,
                        "stage": "STT",
                        "error": "녹취 내용 없음",
                    })

            except Exception as error:
                print(f"    STT 오류: {error}")
                errors.append({
                    "filename": path.name,
                    "stage": "STT",
                    "error": str(error),
                })

            if index % 10 == 0 or index == len(missing_stt):
                print_eta(stt_phase_start, index, len(missing_stt), "STT")

        # CTranslate2/CUDA 모델을 del/gc 하는 순간 일부 Windows 환경에서
        # 프로세스가 조용히 종료되는 문제가 있어 명시적 해제를 하지 않는다.
        # 새 Python 프로세스로 교체하면 OS가 CUDA 메모리를 확실히 회수한다.
        print("\nSTT 완료. CUDA 메모리 해제를 위해 AI 분석 모드로 재시작합니다...")
        sys.stdout.flush()
        script_path = str(Path(__file__).resolve())
        os.execv(sys.executable, [sys.executable, script_path, "--ai-only"])

    # STT 캐시를 읽어서 AI 입력 목록 구성
    for path in audio_files:
        try:
            metadata = parse_filename(path)
            transcript_file = TRANSCRIPT_DIR / f"{get_cache_key(path)}.txt"

            if not transcript_file.exists():
                errors.append({
                    "filename": path.name,
                    "stage": "STT",
                    "error": "STT 캐시 파일 없음",
                })
                continue

            transcript = transcript_file.read_text(encoding="utf-8")
            if not transcript.strip():
                errors.append({
                    "filename": path.name,
                    "stage": "STT",
                    "error": "녹취 내용 없음",
                })
                continue

            transcripts.append({
                "path": path,
                "filename": path.name,
                "metadata": metadata,
                "transcript": transcript,
            })

        except Exception as error:
            errors.append({
                "filename": path.name,
                "stage": "STT 캐시 읽기",
                "error": str(error),
            })

    # -----------------------------------------------------
    # 2단계: Ollama 분석
    # -----------------------------------------------------

    print("\n[2단계] AI 문의 분석")
    analyze_phase_start = time.perf_counter()
    call_results = []
    analyze_total = len(transcripts)

    for index, call in enumerate(transcripts, start=1):
        file_start = time.perf_counter()
        print(f"[{index}/{analyze_total}] {call['filename']}")

        try:
            analysis, cached = analyze_transcript(call["path"], call["transcript"])
            elapsed = time.perf_counter() - file_start

            call_results.append({
                "filename": call["filename"],
                "metadata": call["metadata"],
                "transcript": call["transcript"],
                "analysis": analysis,
            })

            inquiry_count = len(analysis.get("inquiries", []))
            review_count = sum(
                1 for x in analysis.get("inquiries", [])
                if x.get("review") == "Y"
            )

            cache_text = "캐시" if cached else "신규"
            print(
                f"    AI {cache_text} 완료: {format_duration(elapsed)} | "
                f"문의 {inquiry_count}건 | 검토 {review_count}건"
            )

        except Exception as error:
            elapsed = time.perf_counter() - file_start
            print(f"    AI 오류: {error} ({format_duration(elapsed)})")
            errors.append({
                "filename": call["filename"],
                "stage": "AI",
                "error": str(error),
            })

        if index % CHECKPOINT_EVERY == 0:
            print("    중간 Excel 저장...")
            create_excel(call_results, errors)

        if index % 5 == 0 or index == analyze_total:
            print_eta(analyze_phase_start, index, analyze_total, "AI")

    # -----------------------------------------------------
    # 최종 Excel
    # -----------------------------------------------------

    print("\n최종 Excel 생성 중...")
    excel_start = time.perf_counter()
    create_excel(call_results, errors)
    print(f"Excel 생성 완료: {format_duration(time.perf_counter() - excel_start)}")

    total_inquiries = sum(
        len(call["analysis"].get("inquiries", []))
        for call in call_results
    )
    total_reviews = sum(
        1
        for call in call_results
        for item in call["analysis"].get("inquiries", [])
        if item.get("review") == "Y"
    )

    print("\n" + "=" * 72)
    print("완료")
    print(f"결과파일       : {OUTPUT_FILE}")
    print(f"성공 통화      : {len(call_results)}건")
    print(f"추출 문의      : {total_inquiries}건")
    print(f"검토필요       : {total_reviews}건")
    print(f"오류           : {len(errors)}건")
    print(f"종료시간       : {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}")
    print("=" * 72)


if __name__ == "__main__":
    main()
