import json
import re
import hashlib
import urllib.request
from datetime import datetime
from pathlib import Path

from faster_whisper import WhisperModel
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


# =========================================================
# 설정
# =========================================================

CALL_DIR = Path(r"E:\통화녹음")

# small: 빠름 / medium: 좀 더 정확함
WHISPER_MODEL = "medium"

# CPU에서 실행
WHISPER_DEVICE = "cpu"
WHISPER_COMPUTE_TYPE = "int8"

# NVIDIA GPU + CUDA 환경이 제대로 되어 있다면:
# WHISPER_DEVICE = "cuda"
# WHISPER_COMPUTE_TYPE = "float16"

OLLAMA_MODEL = "qwen3:4b"
OLLAMA_URL = "http://localhost:11434/api/chat"

OUTPUT_FILE = CALL_DIR / "통화문의내역.xlsx"
CACHE_DIR = CALL_DIR / "_분석캐시"
TRANSCRIPT_DIR = CACHE_DIR / "녹취록"
ANALYSIS_DIR = CACHE_DIR / "분석결과"

AUDIO_EXTENSIONS = {
    ".m4a", ".mp3", ".wav", ".aac",
    ".flac", ".ogg", ".mp4", ".3gp"
}


# =========================================================
# 폴더 생성
# =========================================================

CACHE_DIR.mkdir(exist_ok=True)
TRANSCRIPT_DIR.mkdir(exist_ok=True)
ANALYSIS_DIR.mkdir(exist_ok=True)


# =========================================================
# 파일명에서 날짜 / 시간 / 상대방 추출
# =========================================================

def parse_filename(path: Path):
    """
    예:
    통화 디앤디기술 백현욱차장님_261007_091215.m4a

    ->
    상대방: 디앤디기술 백현욱차장님
    날짜: 2026-10-07
    시간: 09:12:15
    """

    stem = path.stem

    pattern = r"^(?:통화(?:\s*녹음)?\s*)?(.*?)_(\d{6}|\d{8})_(\d{6})$"

    match = re.match(pattern, stem)

    if match:
        person = match.group(1).strip()
        date_string = match.group(2)
        time_string = match.group(3)

        if len(date_string) == 6:
            date_obj = datetime.strptime(date_string, "%y%m%d")
        else:
            date_obj = datetime.strptime(date_string, "%Y%m%d")

        time_obj = datetime.strptime(time_string, "%H%M%S")

        return {
            "person": person,
            "date": date_obj.strftime("%Y-%m-%d"),
            "time": time_obj.strftime("%H:%M:%S")
        }

    # 파일명이 예상 형식과 다르면 파일 수정날짜 사용
    modified = datetime.fromtimestamp(path.stat().st_mtime)

    return {
        "person": stem,
        "date": modified.strftime("%Y-%m-%d"),
        "time": modified.strftime("%H:%M:%S")
    }


# =========================================================
# 캐시 파일 이름
# =========================================================

def get_cache_key(path: Path):
    value = str(path.resolve()).encode("utf-8")
    return hashlib.sha1(value).hexdigest()


# =========================================================
# Whisper 음성 → 텍스트
# =========================================================

def transcribe_audio(path: Path, model: WhisperModel):

    cache_key = get_cache_key(path)
    transcript_file = TRANSCRIPT_DIR / f"{cache_key}.txt"

    # 이미 변환했다면 다시 Whisper 실행하지 않음
    if transcript_file.exists():
        return transcript_file.read_text(
            encoding="utf-8"
        )

    print("    음성 → 텍스트 변환 중...")

    segments, info = model.transcribe(
        str(path),
        language="ko",
        beam_size=5,
        vad_filter=True,
        vad_parameters={
            "min_silence_duration_ms": 500
        }
    )

    lines = []

    for segment in segments:

        text = segment.text.strip()

        if not text:
            continue

        start = int(segment.start)

        minute = start // 60
        second = start % 60

        lines.append(
            f"[{minute:02d}:{second:02d}] {text}"
        )

    transcript = "\n".join(lines)

    transcript_file.write_text(
        transcript,
        encoding="utf-8"
    )

    return transcript


# =========================================================
# Ollama 호출
# =========================================================

def call_ollama(prompt):

    body = {
        "model": OLLAMA_MODEL,
        "stream": False,
        "format": "json",
        "messages": [
            {
                "role": "system",
                "content": """
당신은 IT 시스템 유지보수 업무의 통화 기록을 정리하는 담당자입니다.

통화 녹취를 보고 실제로 문의한 사항만 추출하세요.

규칙:

1. 인사말, 잡담은 제외합니다.
2. 통화 한 건에 문의사항이 여러 개라면 각각 분리합니다.
3. 녹취에 없는 사실을 추측하지 않습니다.
4. 단순 문의라도 빠뜨리지 않습니다.
5. 문의자의 요구사항과 담당자의 답변/처리내용을 구분합니다.
6. 추후 해야 할 일이 있으면 후속조치에 기록합니다.
7. 이미 해결됐으면 상태를 "완료"로 합니다.
8. 추후 수정/배포/확인이 필요하면 "처리예정"으로 합니다.
9. 정확하지 않으면 "확인필요"로 합니다.
10. 문의가 없다면 inquiries를 빈 배열로 반환합니다.

category는 가능한 한 아래 중 선택하세요.

- 장애/오류
- 성능
- 기능문의
- 기능수정
- 출력/문서
- 데이터
- 계정/권한
- 파일/다운로드
- 배포
- 사용방법
- 기타

반드시 JSON만 반환하세요.

형식:

{
  "summary": "통화 전체 요약",
  "inquiries": [
    {
      "category": "분류",
      "title": "문의 제목",
      "inquiry": "문의 내용",
      "response": "답변 또는 처리 내용",
      "follow_up": "후속 조치",
      "status": "완료 또는 처리예정 또는 확인필요"
    }
  ]
}
"""
            },
            {
                "role": "user",
                "content": prompt
            }
        ],
        "options": {
            "temperature": 0.1,
            "num_ctx": 32768
        }
    }

    data = json.dumps(
        body,
        ensure_ascii=False
    ).encode("utf-8")

    request = urllib.request.Request(
        OLLAMA_URL,
        data=data,
        headers={
            "Content-Type": "application/json"
        }
    )

    with urllib.request.urlopen(
            request,
            timeout=600
    ) as response:

        result = json.loads(
            response.read().decode("utf-8")
        )

    content = result["message"]["content"]

    return json.loads(content)


# =========================================================
# 긴 통화 분할
# =========================================================

def split_transcript(text, max_chars=14000):

    if len(text) <= max_chars:
        return [text]

    lines = text.splitlines()

    chunks = []
    current = []

    current_length = 0

    for line in lines:

        if current_length + len(line) > max_chars:

            chunks.append("\n".join(current))

            # 앞뒤 문맥 약간 유지
            current = current[-5:]

            current_length = sum(
                len(x) for x in current
            )

        current.append(line)
        current_length += len(line)

    if current:
        chunks.append("\n".join(current))

    return chunks


# =========================================================
# AI 분석
# =========================================================

def analyze_transcript(path, transcript):

    cache_key = get_cache_key(path)

    analysis_file = ANALYSIS_DIR / f"{cache_key}.json"

    # 이전 분석 결과가 있으면 재사용
    if analysis_file.exists():

        return json.loads(
            analysis_file.read_text(
                encoding="utf-8"
            )
        )

    chunks = split_transcript(transcript)

    all_inquiries = []
    summaries = []

    for index, chunk in enumerate(chunks, start=1):

        print(
            f"    AI 분석 중... "
            f"({index}/{len(chunks)})"
        )

        prompt = f"""
아래는 실제 업무 통화 녹취입니다.

문의사항을 정리하세요.

[녹취 시작]

{chunk}

[녹취 끝]
"""

        result = call_ollama(prompt)

        summary = result.get("summary", "")

        if summary:
            summaries.append(summary)

        inquiries = result.get(
            "inquiries",
            []
        )

        all_inquiries.extend(inquiries)

    # 중복 제거
    unique = []
    seen = set()

    for item in all_inquiries:

        key = (
            item.get("title", "").strip(),
            item.get("inquiry", "").strip()
        )

        if key not in seen:
            seen.add(key)
            unique.append(item)

    final_result = {
        "summary": " / ".join(summaries),
        "inquiries": unique
    }

    analysis_file.write_text(
        json.dumps(
            final_result,
            ensure_ascii=False,
            indent=2
        ),
        encoding="utf-8"
    )

    return final_result


# =========================================================
# Excel
# =========================================================

def create_excel(call_results):

    wb = Workbook()

    # -----------------------------------------------------
    # 문의내역
    # -----------------------------------------------------

    ws = wb.active
    ws.title = "문의내역"

    headers = [
        "번호",
        "통화일자",
        "통화시간",
        "상대방",
        "분류",
        "문의제목",
        "문의내용",
        "답변/처리내용",
        "후속조치",
        "상태",
        "통화요약",
        "원본파일"
    ]

    ws.append(headers)

    row_number = 1

    for call in call_results:

        metadata = call["metadata"]
        analysis = call["analysis"]

        inquiries = analysis.get(
            "inquiries",
            []
        )

        if not inquiries:
            continue

        for inquiry in inquiries:

            ws.append([
                row_number,
                metadata["date"],
                metadata["time"],
                metadata["person"],
                inquiry.get(
                    "category",
                    ""
                ),
                inquiry.get(
                    "title",
                    ""
                ),
                inquiry.get(
                    "inquiry",
                    ""
                ),
                inquiry.get(
                    "response",
                    ""
                ),
                inquiry.get(
                    "follow_up",
                    ""
                ),
                inquiry.get(
                    "status",
                    ""
                ),
                analysis.get(
                    "summary",
                    ""
                ),
                call["filename"]
            ])

            row_number += 1

    # -----------------------------------------------------
    # 통화목록
    # -----------------------------------------------------

    ws_calls = wb.create_sheet(
        "통화목록"
    )

    call_headers = [
        "번호",
        "통화일자",
        "통화시간",
        "상대방",
        "문의건수",
        "통화요약",
        "원본파일"
    ]

    ws_calls.append(call_headers)

    for index, call in enumerate(
            call_results,
            start=1
    ):

        metadata = call["metadata"]
        analysis = call["analysis"]

        ws_calls.append([
            index,
            metadata["date"],
            metadata["time"],
            metadata["person"],
            len(
                analysis.get(
                    "inquiries",
                    []
                )
            ),
            analysis.get(
                "summary",
                ""
            ),
            call["filename"]
        ])

    # -----------------------------------------------------
    # 스타일
    # -----------------------------------------------------

    for sheet in [ws, ws_calls]:

        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions

        for cell in sheet[1]:

            cell.font = Font(
                bold=True
            )

            cell.fill = PatternFill(
                "solid",
                fgColor="D9EAF7"
            )

            cell.alignment = Alignment(
                horizontal="center",
                vertical="center"
            )

        for row in sheet.iter_rows(
                min_row=2
        ):

            for cell in row:

                cell.alignment = Alignment(
                    vertical="top",
                    wrap_text=True
                )

    widths = {
        1: 7,
        2: 13,
        3: 11,
        4: 28,
        5: 14,
        6: 30,
        7: 55,
        8: 55,
        9: 45,
        10: 13,
        11: 60,
        12: 45
    }

    for col, width in widths.items():

        ws.column_dimensions[
            get_column_letter(col)
        ].width = width

    widths_calls = {
        1: 7,
        2: 13,
        3: 11,
        4: 28,
        5: 10,
        6: 70,
        7: 45
    }

    for col, width in widths_calls.items():

        ws_calls.column_dimensions[
            get_column_letter(col)
        ].width = width

    wb.save(OUTPUT_FILE)


# =========================================================
# Ollama 실행 확인
# =========================================================

def check_ollama():

    try:

        with urllib.request.urlopen(
                "http://localhost:11434/api/tags",
                timeout=5
        ) as response:

            return response.status == 200

    except Exception:
        return False


# =========================================================
# MAIN
# =========================================================

def main():

    print("=" * 60)
    print("갤럭시 통화녹음 업무문의 자동 분석")
    print("=" * 60)

    if not CALL_DIR.exists():

        print(
            f"폴더가 없습니다: {CALL_DIR}"
        )

        return

    if not check_ollama():

        print()
        print("Ollama에 연결할 수 없습니다.")
        print()
        print("Ollama를 실행한 뒤:")
        print()
        print(
            f"ollama pull {OLLAMA_MODEL}"
        )
        print()

        return

    audio_files = sorted(
        [
            file
            for file in CALL_DIR.iterdir()
            if file.is_file()
               and file.suffix.lower()
               in AUDIO_EXTENSIONS
        ]
    )

    print()
    print(
        f"발견된 통화녹음: "
        f"{len(audio_files)}개"
    )
    print()

    if not audio_files:
        return

    print(
        f"Whisper 모델 로딩: "
        f"{WHISPER_MODEL}"
    )

    whisper = WhisperModel(
        WHISPER_MODEL,
        device=WHISPER_DEVICE,
        compute_type=WHISPER_COMPUTE_TYPE
    )

    results = []

    for index, path in enumerate(
            audio_files,
            start=1
    ):

        print()
        print(
            f"[{index}/{len(audio_files)}] "
            f"{path.name}"
        )

        try:

            metadata = parse_filename(
                path
            )

            transcript = transcribe_audio(
                path,
                whisper
            )

            if not transcript.strip():

                print(
                    "    녹취 내용 없음 - 건너뜀"
                )

                continue

            analysis = analyze_transcript(
                path,
                transcript
            )

            results.append({
                "filename": path.name,
                "metadata": metadata,
                "transcript": transcript,
                "analysis": analysis
            })

            count = len(
                analysis.get(
                    "inquiries",
                    []
                )
            )

            print(
                f"    완료 - 문의 {count}건"
            )

        except Exception as error:

            print(
                f"    오류: {error}"
            )

    print()
    print("Excel 생성 중...")

    create_excel(results)

    print()
    print("=" * 60)
    print("완료")
    print(f"결과: {OUTPUT_FILE}")
    print("=" * 60)


if __name__ == "__main__":
    main()